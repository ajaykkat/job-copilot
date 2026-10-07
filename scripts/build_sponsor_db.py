"""
Build data/sponsors.db from DOL LCA disclosure files.

1. Download one or more "LCA Programs (H-1B, H-1B1, E-3)" disclosure files from
   https://www.dol.gov/agencies/eta/foreign-labor/performance
   (e.g. LCA_Disclosure_Data_FY2024_Q4.xlsx). The last 3 fiscal years is plenty.
2. Run:
       python scripts/build_sponsor_db.py ~/Downloads/LCA_Disclosure_Data_FY2024_Q4.xlsx ...

Accepts .xlsx and .csv. Rebuilds the database from scratch each run.
"""

import argparse
import csv
import os
import re
import sqlite3
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from sponsorship import DB_PATH, annualize, normalize_company  # noqa: E402

SCHEMA = """
CREATE TABLE filings (
    employer_norm TEXT NOT NULL,
    employer      TEXT,
    job_title     TEXT,
    soc_title     TEXT,
    wage          REAL,
    pw            REAL,
    pw_level      TEXT,
    state         TEXT,
    fy            INTEGER,
    status        TEXT,
    visa_class    TEXT
);
"""
INDEX = "CREATE INDEX idx_employer_norm ON filings(employer_norm);"

# Column names in the DOL files (FY2020+); older files used slightly different ones
COLUMNS = {
    "status":     ["CASE_STATUS"],
    "employer":   ["EMPLOYER_NAME"],
    "job_title":  ["JOB_TITLE"],
    "soc_title":  ["SOC_TITLE", "SOC_NAME"],
    "wage":       ["WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY"],
    "wage_unit":  ["WAGE_UNIT_OF_PAY"],
    "pw":         ["PREVAILING_WAGE"],
    "pw_unit":    ["PW_UNIT_OF_PAY"],
    "pw_level":   ["PW_WAGE_LEVEL"],
    "state":      ["WORKSITE_STATE", "WORKSITE_STATE_1"],
    "decision":   ["DECISION_DATE"],
    "visa_class": ["VISA_CLASS"],
}


def _rows_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().upper() for h in next(it)]
    for row in it:
        yield dict(zip(header, row))
    wb.close()


def _rows_csv(path):
    with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
        reader = csv.reader(f)
        header = [h.strip().upper() for h in next(reader)]
        for row in reader:
            yield dict(zip(header, row))


def _pick(row: dict, key: str):
    for col in COLUMNS[key]:
        if col in row and row[col] not in (None, ""):
            return row[col]
    return None


def _fiscal_year(path: str, decision) -> int | None:
    m = re.search(r"FY\s*(\d{4})", os.path.basename(path), re.I)
    if m:
        return int(m.group(1))
    # Fiscal year starts Oct 1
    s = str(decision or "")
    m = re.match(r"(\d{4})-(\d{2})", s) or re.match(r"(\d{1,2})/\d{1,2}/(\d{4})", s)
    if not m:
        return None
    if "-" in s[:5]:
        year, month = int(m.group(1)), int(m.group(2))
    else:
        month, year = int(m.group(1)), int(m.group(2))
    return year + 1 if month >= 10 else year


def import_file(conn, path: str) -> int:
    rows = _rows_xlsx(path) if path.lower().endswith((".xlsx", ".xlsm")) else _rows_csv(path)
    batch, count = [], 0
    for row in rows:
        employer = _pick(row, "employer")
        norm     = normalize_company(str(employer or ""))
        if not norm:
            continue
        batch.append((
            norm,
            str(employer).strip(),
            str(_pick(row, "job_title") or "").strip(),
            str(_pick(row, "soc_title") or "").strip(),
            annualize(_pick(row, "wage"), str(_pick(row, "wage_unit") or "Year")),
            annualize(_pick(row, "pw"),   str(_pick(row, "pw_unit")   or "Year")),
            str(_pick(row, "pw_level") or "").strip(),
            str(_pick(row, "state") or "").strip().upper()[:2],
            _fiscal_year(path, _pick(row, "decision")),
            str(_pick(row, "status") or "").strip().upper(),
            str(_pick(row, "visa_class") or "").strip(),
        ))
        if len(batch) >= 5000:
            conn.executemany("INSERT INTO filings VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
            count += len(batch)
            batch.clear()
    if batch:
        conn.executemany("INSERT INTO filings VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
        count += len(batch)
    return count


def build(paths: list, out: str = DB_PATH) -> int:
    os.makedirs(os.path.dirname(out), exist_ok=True)
    tmp = out + ".tmp"
    if os.path.exists(tmp):
        os.remove(tmp)
    conn = sqlite3.connect(tmp)
    conn.executescript(SCHEMA)
    total = 0
    for p in paths:
        n = import_file(conn, p)
        print(f"  {os.path.basename(p)}: {n:,} rows")
        total += n
    conn.executescript(INDEX)
    conn.commit()
    conn.close()
    os.replace(tmp, out)   # atomic swap — the running app never sees a half-built DB
    return total


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", help="LCA disclosure .xlsx or .csv files")
    ap.add_argument("--out", default=DB_PATH, help=f"output SQLite path (default {DB_PATH})")
    args = ap.parse_args()
    total = build(args.files, args.out)
    print(f"\n  Done — {total:,} filings → {args.out}")


if __name__ == "__main__":
    main()
