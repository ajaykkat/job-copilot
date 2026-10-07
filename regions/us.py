"""
United States — H-1B history from DOL LCA disclosure data.

Download "LCA Programs (H-1B, H-1B1, E-3)" .xlsx files from
https://www.dol.gov/agencies/eta/foreign-labor/performance
and import them with: python scripts/build_sponsor_db.py --country US <files>
"""

import csv
import os
import re
import sqlite3
import statistics
from collections import Counter

from regions.base import DataRegion, money, normalize_company

_TITLE_NOISE = {
    "senior", "sr", "junior", "jr", "lead", "staff", "principal", "associate",
    "i", "ii", "iii", "iv", "1", "2", "3", "4", "level", "entry", "mid",
    "remote", "hybrid", "onsite", "contract", "fulltime", "full", "time",
    "the", "a", "an", "of", "and", "for", "with", "to", "in", "at",
}

_ACRONYMS = {"qa", "ui", "ux", "api", "aws", "sap", "etl", "it", "sre", "ml", "ai", "bi",
             "dba", "sql", "erp", "crm", "hr", "vp", "cto", "ios", "net", "ii", "iii", "iv"}

_WAGE_MULT = {"year": 1, "hour": 2080, "week": 52, "bi-weekly": 26, "biweekly": 26, "month": 12}

US_STATES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID", "IL", "IN", "IA",
    "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ",
    "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT",
    "VA", "WA", "WV", "WI", "WY", "DC",
}


def title_tokens(title: str) -> set:
    s = re.sub(r"[^a-z0-9+#]+", " ", (title or "").lower())
    return {w for w in s.split() if w not in _TITLE_NOISE}


def display_title(title: str) -> str:
    """Readable job title: DOL rows are often ALL CAPS ("SR. QA ANALYST")."""
    t = " ".join((title or "").split())
    if not t.isupper():
        return t
    words = []
    for w in t.split(" "):
        core = re.sub(r"[^A-Za-z]", "", w).lower()
        words.append(w if core in _ACRONYMS else w.capitalize())
    return " ".join(words)


def annualize(amount, unit: str):
    """Convert a wage to an annual figure; None if unusable."""
    try:
        val = float(str(amount).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None
    mult = _WAGE_MULT.get((unit or "year").strip().lower())
    if not mult or val <= 0:
        return None
    annual = val * mult
    # Guard against obviously broken rows (hourly rate tagged as yearly, etc.)
    return annual if 15_000 <= annual <= 1_500_000 else None


# ── Import ───────────────────────────────────────────────────────────────────

SCHEMA = """
CREATE TABLE filings (
    employer_norm TEXT NOT NULL, employer TEXT, job_title TEXT, soc_title TEXT,
    wage REAL, pw REAL, pw_level TEXT, state TEXT, fy INTEGER, status TEXT, visa_class TEXT
);
"""
INDEX = "CREATE INDEX idx_employer_norm ON filings(employer_norm);"

# Column names in the DOL files (FY2020+); older files used slightly different ones
COLUMNS = {
    "status":     ["CASE_STATUS"],
    "employer":   ["EMPLOYER_NAME"],
    "job_title":  ["JOB_TITLE"],
    "soc_title":  ["SOC_TITLE", "SOC_NAME"],
    "wage":       ["WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY"],
    "wage_unit":  ["WAGE_UNIT_OF_PAY"],
    "pw":         ["PREVAILING_WAGE"],
    "pw_unit":    ["PW_UNIT_OF_PAY"],
    "pw_level":   ["PW_WAGE_LEVEL"],
    "state":      ["WORKSITE_STATE", "WORKSITE_STATE_1"],
    "decision":   ["DECISION_DATE"],
    "visa_class": ["VISA_CLASS"],
}


def read_rows(path: str):
    """Yield {UPPERCASE_HEADER: value} dicts from an .xlsx or .csv file."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        header = [str(h or "").strip().upper() for h in next(it)]
        for row in it:
            yield dict(zip(header, row))
        wb.close()
    else:
        with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
            reader = csv.reader(f)
            header = [h.strip().upper() for h in next(reader)]
            for row in reader:
                yield dict(zip(header, row))


def _pick(row: dict, key: str):
    for col in COLUMNS[key]:
        if col in row and row[col] not in (None, ""):
            return row[col]
    return None


def fiscal_year(path: str, decision) -> int | None:
    m = re.search(r"FY\s*(\d{4})", os.path.basename(path), re.I)
    if m:
        return int(m.group(1))
    # US federal fiscal year starts Oct 1
    s = str(decision or "")
    m = re.match(r"(\d{4})-(\d{2})", s)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
    else:
        m = re.match(r"(\d{1,2})/\d{1,2}/(\d{4})", s)
        if not m:
            return None
        month, year = int(m.group(1)), int(m.group(2))
    return year + 1 if month >= 10 else year


# ── Region ───────────────────────────────────────────────────────────────────

class UnitedStates(DataRegion):
    code, name, currency = "US", "United States", "USD"
    aliases = ("united states", "united states of america", "usa", "u.s.")
    data_source = "US H-1B filings (DOL)"
    data_setup  = {
        "url": "https://www.dol.gov/agencies/eta/foreign-labor/performance",
        "file_hint": "the latest “LCA Programs (H-1B, H-1B1, E-3)” .xlsx files (last 2–3 fiscal years)",
    }
    table = "filings"

    def _status_detail(self, conn) -> str:
        lo, hi = conn.execute("SELECT MIN(fy), MAX(fy) FROM filings").fetchone()
        if not lo:
            return ""
        return f"FY{lo}" if lo == hi else f"FY{lo}–{hi}"

    def authorized_by_status(self, profile: dict):
        from user_profile import get_immigration_answers
        if not (profile or {}).get("immigration", {}).get("status"):
            return None
        imm = get_immigration_answers(profile)
        return imm.get("sponsor_now") == "No" and imm.get("sponsor_future") == "No"

    def lookup(self, company: str, title: str = "") -> dict | None:
        norm = normalize_company(company)
        if not norm or not self.available:
            return None
        cols = "employer, job_title, wage, pw, fy, state"
        try:
            rows = self._match(
                f"SELECT {cols} FROM filings WHERE employer_norm = ? AND status LIKE 'CERTIFIED%'",
                f"SELECT {cols} FROM filings WHERE employer_norm LIKE ? AND status LIKE 'CERTIFIED%'",
                norm,
            )
        except sqlite3.Error:
            return None
        if not rows:
            return None

        by_year = Counter(r[4] for r in rows if r[4])
        names   = Counter(r[0] for r in rows)
        wages   = [r[2] for r in rows if r[2]]
        titles  = Counter(display_title(r[1]) for r in rows if r[1])
        last_fy = max(by_year) if by_year else None

        n = len(rows)
        reasons = [f"{n:,} certified H-1B filings" + (f" (last FY{last_fy})" if last_fy else "")]
        stats   = [{"value": f"{n:,}", "label": "H-1B filings"}]

        role_hits, role_pay, role_pw = 0, None, None
        want = title_tokens(title)
        if want:
            hits = [r for r in rows if want <= title_tokens(r[1])]
            if hits:
                role_hits = len(hits)
                hw  = [r[2] for r in hits if r[2]]
                hpw = [r[3] for r in hits if r[3]]
                role_pay = statistics.median(hw) if hw else None
                role_pw  = statistics.median(hpw) if hpw else None
                reasons.append(f"{role_hits:,} for a matching role")
                stats.append({"value": f"{role_hits:,}", "label": "for this role"})

        pay = role_pay or (statistics.median(wages) if wages else None)
        stats.append({"value": money(pay, self.currency), "label": "median pay, role" if role_pay else "median pay"})
        if role_pw:
            stats.append({"value": money(role_pw, self.currency), "label": "prevailing wage"})

        return {
            "strength": "strong" if role_hits or n >= 10 else "some",
            "names":    [nm for nm, _ in names.most_common(3)],
            "reasons":  reasons,
            "stats":    stats,
            "chart":    {"title": "Filings by fiscal year",
                         "bars": [[f"FY{str(fy)[2:]}", c] for fy, c in sorted(by_year.items())]},
            "table":    {"title": "Most-filed roles", "rows": [list(t) for t in titles.most_common(5)]},
            "summary":  f"{n:,} H-1B filings",
            "pay":      money(pay, self.currency) if pay else None,
        }

    def build(self, paths: list, out: str) -> int:
        conn = sqlite3.connect(out)
        conn.executescript(SCHEMA)
        total = 0
        for path in paths:
            batch, count = [], 0
            for row in read_rows(path):
                employer = _pick(row, "employer")
                norm = normalize_company(str(employer or ""))
                if not norm:
                    continue
                batch.append((
                    norm, str(employer).strip(),
                    str(_pick(row, "job_title") or "").strip(),
                    str(_pick(row, "soc_title") or "").strip(),
                    annualize(_pick(row, "wage"), str(_pick(row, "wage_unit") or "Year")),
                    annualize(_pick(row, "pw"),   str(_pick(row, "pw_unit")   or "Year")),
                    str(_pick(row, "pw_level") or "").strip(),
                    str(_pick(row, "state") or "").strip().upper()[:2],
                    fiscal_year(path, _pick(row, "decision")),
                    str(_pick(row, "status") or "").strip().upper(),
                    str(_pick(row, "visa_class") or "").strip(),
                ))
                if len(batch) >= 5000:
                    conn.executemany("INSERT INTO filings VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
                    count += len(batch)
                    batch.clear()
            if batch:
                conn.executemany("INSERT INTO filings VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
                count += len(batch)
            print(f"  {os.path.basename(path)}: {count:,} rows")
            total += count
        conn.executescript(INDEX)
        conn.commit()
        conn.close()
        return total
